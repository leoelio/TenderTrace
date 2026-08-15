from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from tendertrace.adapters.ccgp import Attachment, Notice
from tendertrace.config import Settings
from tendertrace.pipeline.attachments import enrich_attachment_snapshots


class PipelineAttachmentsTests(unittest.TestCase):
    def test_private_network_attachment_is_skipped_before_download(self) -> None:
        def downloader(url: str, max_bytes: int) -> bytes:
            raise AssertionError("unsafe URL must not reach downloader")

        notice = Notice(
            id="internal",
            source_site="partner",
            title="内部地址附件",
            publish_time="2026-08-16",
            region="上海",
            purchaser="测试单位",
            source_url="https://example.com/notice",
            attachments=[
                Attachment(name="internal", url="http://127.0.0.1/private.pdf")
            ],
        )

        with tempfile.TemporaryDirectory() as tmp:
            result = enrich_attachment_snapshots(
                [notice],
                settings=Settings.load(Path(tmp)),
                downloader=downloader,
            )

        record = result.notices[0].fields["attachment_snapshots"][0]
        self.assertEqual(record["status"], "skipped")
        self.assertEqual(record["error"], "unsafe_url")

    def test_enrich_attachment_snapshots_downloads_extracts_and_limits(self) -> None:
        files = {
            "https://example.com/spec.pdf": _pdf_bytes("project SH-2026-001 budget 120"),
            "https://example.com/spec.docx": _docx_bytes("docx project requirement server"),
            "https://example.com/spec.xlsx": _xlsx_bytes("xlsx equipment list"),
            "https://example.com/extra.pdf": _pdf_bytes("extra should be skipped"),
        }

        def fake_download(url: str, max_bytes: int) -> bytes:
            self.assertGreater(max_bytes, 100)
            return files[url]

        notice = Notice(
            id="n1",
            source_site="ccgp",
            title="上海设备采购公告",
            publish_time="2026-07-06",
            region="上海",
            purchaser="上海某单位",
            source_url="https://www.ccgp.gov.cn/detail.htm",
            attachments=[
                Attachment(name="PDF采购需求", url="https://example.com/spec.pdf"),
                Attachment(name="DOCX采购需求", url="https://example.com/spec.docx"),
                Attachment(name="XLSX清单", url="https://example.com/spec.xlsx"),
                Attachment(name="额外附件", url="https://example.com/extra.pdf"),
            ],
        )

        with tempfile.TemporaryDirectory() as tmp:
            settings = Settings.load(Path(tmp))
            result = enrich_attachment_snapshots(
                [notice],
                settings=settings,
                max_per_notice=3,
                max_bytes=1024 * 1024,
                downloader=fake_download,
            )

            records = result.notices[0].fields["attachment_snapshots"]
            stored_files = list((settings.snapshots_dir / "attachments").glob("*/*"))

        self.assertEqual(result.stats["attachments_seen"], 4)
        self.assertEqual(result.stats["attachments_attempted"], 3)
        self.assertEqual(result.stats["attachments_downloaded"], 3)
        self.assertEqual(result.stats["attachments_extracted"], 3)
        self.assertEqual(result.stats["attachments_skipped"], 1)
        self.assertEqual(records[0]["status"], "extracted")
        self.assertIn("SH-2026-001", records[0]["text_excerpt"])
        self.assertEqual(records[1]["status"], "extracted")
        self.assertIn("docx project requirement", records[1]["text_excerpt"])
        self.assertEqual(records[2]["status"], "extracted")
        self.assertIn("xlsx equipment list", records[2]["text_excerpt"])
        self.assertEqual(records[3]["status"], "skipped")
        self.assertTrue(stored_files)

    def test_unsupported_attachment_type_is_skipped_without_download(self) -> None:
        def fail_if_called(url: str, max_bytes: int) -> bytes:
            raise AssertionError("unsupported attachment should not be downloaded")

        notice = Notice(
            id="n2",
            source_site="ccgp",
            title="上海设备采购公告",
            publish_time="2026-07-06",
            region="上海",
            purchaser="上海某单位",
            source_url="https://www.ccgp.gov.cn/detail.htm",
            attachments=[Attachment(name="压缩包", url="https://example.com/files.zip")],
        )

        with tempfile.TemporaryDirectory() as tmp:
            settings = Settings.load(Path(tmp))
            result = enrich_attachment_snapshots(
                [notice],
                settings=settings,
                downloader=fail_if_called,
            )

        record = result.notices[0].fields["attachment_snapshots"][0]
        self.assertEqual(result.stats["attachments_skipped"], 1)
        self.assertEqual(record["status"], "skipped")
        self.assertEqual(record["error"], "unsupported_text_extraction")


def _pdf_bytes(text: str) -> bytes:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 720, text)
    pdf.save()
    return buffer.getvalue()


def _docx_bytes(text: str) -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(text: str) -> bytes:
    buffer = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "需求"
    sheet.append([text])
    workbook.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    unittest.main()
